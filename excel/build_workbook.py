# Builds Off-Grid-Solar-Sizing-Workbook.xlsx with openpyxl. Formulas only; mirrors src/engine.js.
# Run: python excel/build_workbook.py
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Off-Grid-Solar-Sizing-Workbook.xlsx")

SQ = "²"; DIV = "÷"; ARROW = "→"; TIMES = "×"; DASH = "–"

INVERTERS = [300, 500, 1000, 1500, 2000, 3000, 4000, 5000, 6000, 8000, 10000]
CONTROLLERS = [10, 20, 30, 40, 50, 60, 80, 100]
CABLES = [(1.5, 15), (2.5, 21), (4, 28), (6, 36), (10, 50), (16, 66), (25, 84), (35, 104), (50, 125), (70, 160), (95, 194), (120, 225)]
FUSES = [10, 15, 20, 25, 30, 40, 50, 63, 80, 100, 125, 160, 200, 250, 300, 400]
MCBS = [6, 10, 16, 20, 25, 32, 40, 50, 63]

# ---- styles ----
TITLE = Font(name="Calibri", size=16, bold=True, color="1F4E78")
H2 = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
H2_FILL = PatternFill("solid", fgColor="1F4E78")
BOLD = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
INPUT_FONT = Font(color="0000CC")
CALC_FILL = PatternFill("solid", fgColor="EDEDED")
WARN_FONT = Font(color="C00000", bold=True)
NOTE_FONT = Font(italic=True, color="595959", size=9)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
names = {}


def define(name, ref):
    names[name] = ref
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def header(ws, row, text, ncols=4):
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = H2_FILL
    ws.cell(row=row, column=1, value=text).font = H2


def input_cell(cell):
    cell.fill = INPUT_FILL; cell.font = INPUT_FONT; cell.border = BORDER
    cell.protection = cell.protection.copy(locked=False)


def calc_cell(cell):
    cell.fill = CALC_FILL; cell.border = BORDER


# =====================================================================
# Tables
# =====================================================================
wt = wb.active; wt.title = "Tables"
wt["A1"] = "Standard lookup tables (edit only if your market uses different standard sizes; keep lists sorted ascending)"
wt["A1"].font = NOTE_FONT
cols = [("A", "Inverter W", INVERTERS, "InverterSizes"), ("C", "MPPT controller A", CONTROLLERS, "ControllerSizes"),
        ("E", "DC fuse A", FUSES, "FuseSizes"), ("G", "AC MCB A", MCBS, "MCBSizes")]
for col, title, vals, nm in cols:
    wt[f"{col}2"] = title; wt[f"{col}2"].font = BOLD
    for i, v in enumerate(vals):
        c = wt[f"{col}{3 + i}"]; c.value = v; input_cell(c)
    define(nm, f"Tables!${col}$3:${col}${2 + len(vals)}")
wt["I2"] = "Cable mm" + SQ; wt["J2"] = "Ampacity A"; wt["I2"].font = BOLD; wt["J2"].font = BOLD
for i, (a, amp) in enumerate(CABLES):
    c = wt[f"I{3 + i}"]; c.value = a; input_cell(c)
    c = wt[f"J{3 + i}"]; c.value = amp; input_cell(c)
define("CableArea", f"Tables!$I$3:$I${2 + len(CABLES)}")
define("CableAmp", f"Tables!$J$3:$J${2 + len(CABLES)}")
wt["L2"] = "Copper resistivity (ohm mm" + SQ + "/m)"; wt["L2"].font = BOLD
wt["L3"] = 0.0175; input_cell(wt["L3"]); define("CopperResistivity", "Tables!$L$3")
wt["I16"] = ("Ampacities are indicative values for single-core PVC copper cable in free air at about 30 C. "
             "Derate for conduit, heat and bundling.")
wt["I16"].font = NOTE_FONT
for col, w in zip("ABCDEFGHIJKL", [12, 3, 18, 3, 12, 3, 12, 3, 12, 12, 3, 30]):
    wt.column_dimensions[col].width = w

# =====================================================================
# Inputs
# =====================================================================
wi = wb.create_sheet("Inputs", 0)
LF, LL = 8, 32  # load rows
wi["A1"] = "Off-Grid Solar Sizing Workbook " + DASH + " Inputs"; wi["A1"].font = TITLE
wi["A2"] = "Yellow cells are inputs, grey cells are formulas. Results update automatically; see the Results and Report sheets."
wi["A2"].font = NOTE_FONT
for r, (lab, nm, default) in enumerate([("Project", "ProjectName", "Rural clinic (example)"),
                                       ("Prepared by", "PreparedBy", ""), ("Date", "ReportDate", "")], start=3):
    wi.cell(row=r, column=1, value=lab).font = BOLD
    c = wi.cell(row=r, column=2, value=default); input_cell(c); define(nm, f"Inputs!$B${r}")
wi["B5"].number_format = "yyyy-mm-dd"
wi.merge_cells("B3:D3"); wi.merge_cells("B4:D4")

header(wi, 7, "Loads (up to 25 rows)", 9)
heads = ["Appliance", "Watts", "Qty", "Hours/day", "Surge x", "DC?", "Wh/day", "AC continuous W", "AC surge extra W"]
for i, h in enumerate(heads, start=1):
    c = wi.cell(row=7, column=i, value=h); c.font = H2
example = [("LED light", 10, 5, 6, None, "No"), ("Fridge", 100, 1, 10, 5, "No"), ("Laptop", 65, 1, 4, None, "No")]
for r in range(LF, LL + 1):
    for col in range(1, 7):
        input_cell(wi.cell(row=r, column=col))
    wh = wi.cell(row=r, column=7, value=f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),ISNUMBER(D{r}),B{r}>0,C{r}>0,D{r}>0),B{r}*C{r}*D{r},0)')
    ac = wi.cell(row=r, column=8, value=f'=IF(AND(G{r}>0,F{r}<>"Yes"),B{r}*C{r},0)')
    sg = wi.cell(row=r, column=9, value=f'=IF(AND(G{r}>0,F{r}<>"Yes"),B{r}*C{r}*(IF(AND(ISNUMBER(E{r}),E{r}>0),E{r},1)-1),0)')
    for c in (wh, ac, sg):
        calc_cell(c); c.number_format = "#,##0.0"
for i, row_ in enumerate(example):
    for j, v in enumerate(row_, start=1):
        if v is not None:
            wi.cell(row=LF + i, column=j, value=v)
define("LoadName", f"Inputs!$A${LF}:$A${LL}"); define("LoadWatts", f"Inputs!$B${LF}:$B${LL}")
define("LoadQty", f"Inputs!$C${LF}:$C${LL}"); define("LoadHours", f"Inputs!$D${LF}:$D${LL}")
define("LoadSurge", f"Inputs!$E${LF}:$E${LL}"); define("LoadDC", f"Inputs!$F${LF}:$F${LL}")
define("LoadWh", f"Inputs!$G${LF}:$G${LL}"); define("LoadAcW", f"Inputs!$H${LF}:$H${LL}"); define("LoadSurgeExtra", f"Inputs!$I${LF}:$I${LL}")
wi.cell(row=LL + 1, column=1, value="Surge x: start-up multiple for motors (fridge/pump 3-5). Leave blank for 1. DC? = Yes for loads wired straight to the battery (no inverter).").font = NOTE_FONT

dv_yes = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True); dv_yes.add(f"F{LF}:F{LL}")
dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Enter a number of 0 or more.", errorTitle="Invalid"); dv_num.add(f"B{LF}:C{LL}"); dv_num.add(f"E{LF}:E{LL}")
dv_hours = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", allow_blank=True,
                          error="Hours per day must be between 0 and 24.", errorTitle="Invalid"); dv_hours.add(f"D{LF}:D{LL}")
for dv in (dv_yes, dv_num, dv_hours):
    wi.add_data_validation(dv)

# Assumptions block in K:M
AK = 7
for col in range(11, 14):
    wi.cell(row=AK, column=col).fill = H2_FILL
wi.cell(row=AK, column=11, value="Site and equipment assumptions").font = H2
wi.cell(row=AK, column=12, value="Value").font = H2
wi.cell(row=AK, column=13, value="Note").font = H2
assumptions = [
    ("Peak sun hours (worst month)", "PeakSunHours", 5, "h/day; use worst-month data (PVGIS, NASA POWER)", (0.5, 12)),
    ("System derate", "SystemDerate", 0.77, "temperature, dust, wiring, mismatch, controller losses", (0.3, 1)),
    ("Autonomy days", "AutonomyDays", 2, "days of cloudy-weather reserve", (0.1, 10)),
    ("Depth of discharge", "DepthOfDischarge", 0.8, "LiFePO4 ~0.8, lead-acid ~0.5", (0.1, 1)),
    ("Battery round-trip efficiency", "BatteryEfficiency", 0.95, "LiFePO4 ~0.95, lead-acid ~0.85", (0.5, 1)),
    ("Inverter efficiency", "InverterEfficiency", 0.92, "from datasheet", (0.5, 1)),
    ("Inverter safety factor", "InverterSafetyFactor", 1.25, "margin on continuous AC load", (1, 3)),
    ("Inverter surge multiple", "InverterSurgeMultiple", 2, "surge rating vs continuous rating (datasheet)", (1, 5)),
    ("Controller safety factor", "ControllerSafetyFactor", 1.25, "margin on charge current", (1, 3)),
    ("Diversity factor", "DiversityFactor", 1, "share of AC loads running at the same time", (0.1, 1)),
    ("Panel watts", "PanelWatts", 550, "STC rating of one module", (50, 1000)),
    ("Battery unit volts", "BatteryUnitVolts", 12.8, "LiFePO4 12.8, lead-acid 12", (2, 60)),
    ("Battery unit Ah", "BatteryUnitAh", 100, "capacity of one battery unit", (10, 2000)),
    ("System voltage", "SystemVoltsSel", "auto", "auto / 12 / 24 / 48", None),
    ("Max voltage drop %", "MaxVoltageDropPct", 3, "DC cable runs", (0.5, 10)),
    ("Cable: battery to inverter (m, one way)", "CableBatteryInverterM", 2, "one-way length", (0.1, 100)),
    ("Cable: controller to battery (m, one way)", "CableControllerBatteryM", 3, "one-way length", (0.1, 100)),
    ("AC voltage", "AcVolts", 230, "230 or 120", (100, 480)),
]
for i, (lab, nm, val, note, rng) in enumerate(assumptions):
    r = AK + 1 + i
    wi.cell(row=r, column=11, value=lab)
    c = wi.cell(row=r, column=12, value=val); input_cell(c); define(nm, f"Inputs!$L${r}")
    wi.cell(row=r, column=13, value=note).font = NOTE_FONT
    if rng:
        dv = DataValidation(type="decimal", operator="between", formula1=str(rng[0]), formula2=str(rng[1]),
                            error=f"Enter a value between {rng[0]} and {rng[1]}.", errorTitle="Out of range")
        dv.add(f"L{r}"); wi.add_data_validation(dv)
    else:
        dv = DataValidation(type="list", formula1='"auto,12,24,48"', allow_blank=False); dv.add(f"L{r}"); wi.add_data_validation(dv)
for col, w in zip("ABCDEFGHIJKLM", [22, 9, 6, 10, 9, 7, 11, 15, 16, 2, 38, 9, 46]):
    wi.column_dimensions[col].width = w
wi.freeze_panes = "A8"
wi.protection.sheet = True  # no password: Review > Unprotect Sheet to edit layout

# =====================================================================
# Results
# =====================================================================
wr = wb.create_sheet("Results", 1)
wr["A1"] = "Off-Grid Solar Sizing Workbook " + DASH + " Results"; wr["A1"].font = TITLE
wr["A2"] = "All cells on this sheet are formulas. Change inputs on the Inputs sheet."; wr["A2"].font = NOTE_FONT
row = [3]


def pick(lst, minexpr):
    return f'IF(COUNTIF({lst},"<"&{minexpr})+1>ROWS({lst}),"none",INDEX({lst},COUNTIF({lst},"<"&{minexpr})+1))'


def section(text):
    row[0] += 1; header(wr, row[0], text, 4); row[0] += 1


def add(name, label, formula, unit="", fmt="#,##0.0", note=""):
    r = row[0]
    wr.cell(row=r, column=1, value=label)
    c = wr.cell(row=r, column=2, value="=" + formula); calc_cell(c); c.number_format = fmt
    wr.cell(row=r, column=3, value=unit)
    wr.cell(row=r, column=4, value=note).font = NOTE_FONT
    if name: define(name, f"Results!$B${r}")
    row[0] += 1
    return r


section("Input check")
add("ValidCount", "Valid load rows", 'COUNTIF(LoadWh,">0")', "rows", "0")
add("MaxHours", "Highest hours/day entered", "MAX(LoadHours)", "h", "0.0")
add("Status", "Status", 'IF(ValidCount=0,"ERROR: Add at least one load with watts, quantity and hours per day above zero.",IF(MaxHours>24,"ERROR: Hours per day cannot exceed 24.","OK"))', "", "@")

section("1. Daily energy")
add("DcWh", "DC load energy", 'SUMIF(LoadDC,"Yes",LoadWh)', "Wh/day", note="sum of watts x qty x hours for DC rows")
add("AcWh", "AC load energy", "SUM(LoadWh)-DcWh", "Wh/day", note="sum for AC rows (before inverter losses)")
add("RequiredWh", "Required daily energy", "DcWh+AcWh/InverterEfficiency", "Wh/day", note="DC Wh + AC Wh / inverter efficiency")
add("ContinuousW", "Continuous AC load", "SUM(LoadAcW)*DiversityFactor", "W", note="sum of AC watts x qty x diversity factor")
add("AcLoadCount", "AC load rows", 'COUNTIF(LoadAcW,">0")', "rows", "0")

section("2. System voltage")
add("AutoVolts", "Recommended (auto) voltage", "IF(OR(ContinuousW>2500,RequiredWh>5000),48,IF(OR(ContinuousW>800,RequiredWh>2000),24,12))", "V", "0", "12 V up to 800 W / 2 kWh, 24 V up to 2.5 kW / 5 kWh, else 48 V")
add("NominalVolts", "System voltage used", 'IF(SystemVoltsSel="auto",AutoVolts,--SystemVoltsSel)', "V", "0")
add("Series", "Battery units in series", "MAX(1,ROUND(NominalVolts/BatteryUnitVolts,0))", "", "0", "system V / unit V, rounded")
add("BankVolts", "Battery bank voltage", "Series*BatteryUnitVolts", "V", "0.00")

section("3. PV array")
add("PvRequiredW", "PV power required", "RequiredWh/(PeakSunHours*SystemDerate)", "W", note="daily Wh / (peak sun hours x derate)")
add("Panels", "Panels", "ROUNDUP(PvRequiredW/PanelWatts,0)", "", "0", "rounded up to whole panels")
add("PvArrayW", "Array power", "Panels*PanelWatts", "W", "#,##0")
add("DailyYieldWh", "Estimated daily yield", "PvArrayW*PeakSunHours*SystemDerate", "Wh/day")

section("4. Battery bank")
add("BatteryWhRequired", "Bank capacity required", "RequiredWh*AutonomyDays/(DepthOfDischarge*BatteryEfficiency)", "Wh", note="daily Wh x autonomy / (DoD x efficiency)")
add("Parallel", "Strings in parallel", "ROUNDUP(BatteryWhRequired/(BankVolts*BatteryUnitAh),0)", "", "0")
add("BatteryUnits", "Battery units total", "Series*Parallel", "", "0", "series x parallel")
add("BankAh", "Bank capacity", "Parallel*BatteryUnitAh", "Ah", "#,##0")
add("BankWh", "Bank energy", "BankAh*BankVolts", "Wh", "#,##0")
add("UsableWh", "Usable energy", "BankWh*DepthOfDischarge", "Wh", "#,##0")

section("5. Inverter")
add("SurgeExtraW", "Largest start-up surge extra", "IF(AcLoadCount>0,MAX(LoadSurgeExtra),0)", "W", note="largest single load's (surge x - 1) x watts x qty")
add("SurgeW", "Peak surge load", "ContinuousW+SurgeExtraW", "W")
add("MinByContinuous", "Minimum by continuous load", "ContinuousW*InverterSafetyFactor", "W", note="continuous x safety factor")
add("MinBySurge", "Minimum by surge", "SurgeW/InverterSurgeMultiple", "W", note="surge / inverter surge multiple")
add("InverterMinW", "Minimum inverter rating", "MAX(MinByContinuous,MinBySurge)", "W")
add("InverterW", "Inverter rating (standard size)", 'IF(AcLoadCount=0,"n/a",' + pick("InverterSizes", "InverterMinW") + ')', "W", "#,##0", "next standard size; 'none' = above largest listed")
add("LimitedBy", "Limited by", 'IF(AcLoadCount=0,"n/a",IF(MinBySurge>MinByContinuous,"surge","continuous load"))', "", "@")
add("AcBreakerMinA", "AC breaker minimum", "IF(ISNUMBER(InverterW),InverterW/AcVolts*1.25,0)", "A", "0.00", "inverter W / AC V x 1.25")
add("AcBreakerA", "AC output MCB", 'IF(ISNUMBER(InverterW),' + pick("MCBSizes", "AcBreakerMinA") + ',"n/a")', "A", "0")

section("6. MPPT charge controller")
add("ChargeCurrentA", "Charge current (battery side)", "PvArrayW/BankVolts", "A", "0.00", "array W / bank V")
add("ControllerMinA", "Minimum controller rating", "ChargeCurrentA*ControllerSafetyFactor", "A", "0.00")
add("ControllerA", "Controller rating (standard size)", pick("ControllerSizes", "ControllerMinA"), "A", "0")

add("InverterInputA", "Inverter DC input current", "IF(ISNUMBER(InverterW),InverterW/(BankVolts*InverterEfficiency),0)", "A", "0.00", "inverter W / (bank V x inverter efficiency)")


def cable(prefix, label, current_name, length_name, guard):
    section(label)
    add(prefix + "CurrentA", "Design current", current_name, "A", "0.00")
    add(prefix + "AllowedDropV", "Allowed voltage drop", "BankVolts*MaxVoltageDropPct/100", "V", "0.000")
    add(prefix + "MinArea", "Minimum area for voltage drop", f"IF({prefix}AllowedDropV>0,2*{length_name}*{prefix}CurrentA*CopperResistivity/{prefix}AllowedDropV,0)", "mm" + SQ, "0.00", "2 x length x I x resistivity / allowed drop (out and back)")
    add(prefix + "FuseA", "DC fuse (standard size)", f'IF({guard},' + pick("FuseSizes", f"{prefix}CurrentA*1.25") + ',"n/a")', "A", "0", "next standard fuse >= 1.25 x design current")
    add(prefix + "Index", "Cable table row", f'IF(ISNUMBER({prefix}FuseA),MAX(COUNTIF(CableArea,"<"&{prefix}MinArea),COUNTIF(CableAmp,"<"&{prefix}FuseA))+1,0)', "", "0", "first row with area >= drop minimum AND ampacity >= fuse")
    add(prefix + "Area", "Cable size", f'IF(NOT({guard}),"n/a",IF(OR({prefix}Index=0,{prefix}Index>ROWS(CableArea)),"none",INDEX(CableArea,{prefix}Index)))', "mm" + SQ, "0.0")
    add(prefix + "Ampacity", "Cable ampacity", f'IF(ISNUMBER({prefix}Area),INDEX(CableAmp,{prefix}Index),"n/a")', "A", "0")
    add(prefix + "DropPct", "Actual voltage drop", f'IF(ISNUMBER({prefix}Area),2*{length_name}*{prefix}CurrentA*CopperResistivity/{prefix}Area/BankVolts*100,"n/a")', "%", "0.00")


cable("BI_", "7. Cable: battery to inverter", "InverterInputA", "CableBatteryInverterM", "InverterInputA>0")
cable("CB_", "8. Cable: controller to battery", "ChargeCurrentA", "CableControllerBatteryM", "TRUE")

section("Warnings")
w_first = row[0]
warnings = [
    'IF(InverterW="none","Inverter above "&MAX(InverterSizes)&" W needed; use parallel inverters or reduce simultaneous load.","")',
    'IF(ControllerA="none","Charge current "&TEXT(ControllerMinA,"0.0")&" A exceeds "&MAX(ControllerSizes)&" A; split the array across several controllers or use a higher system voltage.","")',
    'IF(BI_Area="none","No standard cable up to "&MAX(CableArea)&" mm' + SQ + ' meets the battery-inverter run; shorten it or raise system voltage.","")',
    'IF(CB_Area="none","No standard cable up to "&MAX(CableArea)&" mm' + SQ + ' meets the controller-battery run; shorten it or raise system voltage.","")',
    'IF(OR(PeakSunHours<2,PeakSunHours>7),"Peak sun hours outside 2-7 is unusual; check the site\'s worst-month solar data.","")',
    'IF(AutonomyDays<1,"Less than one day of autonomy: loads will fail after a cloudy day.","")',
    'IF(InverterInputA>150,"Battery current about "&TEXT(InverterInputA,"0")&" A is high; a higher system voltage would reduce cable size and losses.","")',
    'IF(AND(SystemVoltsSel<>"auto",NominalVolts<AutoVolts),"For this load a "&AutoVolts&" V system is usually more practical.","")',
]
for i, f in enumerate(warnings):
    r = row[0]
    wr.cell(row=r, column=1, value=f"Warning {i + 1}")
    c = wr.cell(row=r, column=2, value="=" + f); calc_cell(c); c.font = WARN_FONT
    wr.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    row[0] += 1
define("WarningList", f"Results!$B${w_first}:$B${row[0] - 1}")
add("WarningCount", "Number of warnings", 'COUNTIF(WarningList,"?*")', "", "0")
for col, w in zip("ABCD", [34, 16, 8, 62]):
    wr.column_dimensions[col].width = w
wr.freeze_panes = "A3"
wr.protection.sheet = True

# =====================================================================
# Report
# =====================================================================
wp = wb.create_sheet("Report", 2)
wp["A1"] = '="Off-Grid Solar Sizing Report ' + DASH + ' "&ProjectName'; wp["A1"].font = TITLE
wp["A2"] = '="Prepared by: "&IF(PreparedBy="","-",PreparedBy)&"   Date: "&IF(ReportDate="","-",TEXT(ReportDate,"yyyy-mm-dd"))&"   Status: "&Status'
wp["A2"].font = NOTE_FONT
r = 4


def prow(label, formula, bold=False):
    global r
    wp.cell(row=r, column=1, value=label).font = BOLD if bold else Font()
    c = wp.cell(row=r, column=2, value=formula); c.alignment = WRAP
    if bold: c.font = BOLD
    r += 1


def psection(text):
    global r
    for col in range(1, 4):
        wp.cell(row=r, column=col).fill = H2_FILL
    wp.cell(row=r, column=1, value=text).font = H2
    r += 1


psection("Summary (bill of materials)")
prow("Solar panels", '=Panels&" ' + TIMES + ' "&PanelWatts&" W = "&TEXT(PvArrayW,"#,##0")&" W array"', True)
prow("Battery bank", '=BatteryUnits&" ' + TIMES + ' "&BatteryUnitVolts&" V / "&BatteryUnitAh&" Ah units ("&Series&"S"&Parallel&"P) = "&TEXT(BankVolts,"0.0")&" V, "&TEXT(BankAh,"#,##0")&" Ah, "&TEXT(BankWh,"#,##0")&" Wh"', True)
prow("Inverter", '=IF(ISNUMBER(InverterW),TEXT(InverterW,"#,##0")&" W continuous (limited by "&LimitedBy&")",IF(InverterW="n/a","not needed (DC-only loads)","none fits - see warnings"))', True)
prow("MPPT charge controller", '=IF(ISNUMBER(ControllerA),ControllerA&" A at "&TEXT(BankVolts,"0.0")&" V","none fits - see warnings")', True)
prow("DC fuses", '="Battery-inverter: "&BI_FuseA&" A;  controller-battery: "&CB_FuseA&" A"', True)
prow("DC cables (copper)", '="Battery-inverter: "&BI_Area&" mm' + SQ + ' ("&CableBatteryInverterM&" m);  controller-battery: "&CB_Area&" mm' + SQ + ' ("&CableControllerBatteryM&" m)"', True)
prow("AC output breaker", '=IF(ISNUMBER(AcBreakerA),AcBreakerA&" A MCB","n/a")', True)
r += 1
psection("Calculation steps")
prow("1. Daily energy", '="AC "&TEXT(AcWh,"#,##0")&" Wh ' + DIV + ' "&InverterEfficiency&" + DC "&TEXT(DcWh,"#,##0")&" Wh = "&TEXT(RequiredWh,"#,##0.0")&" Wh/day   (continuous AC load "&TEXT(ContinuousW,"#,##0")&" W)"')
prow("2. System voltage", '=IF(SystemVoltsSel="auto","Auto: "&TEXT(ContinuousW,"#,##0")&" W continuous and "&TEXT(RequiredWh,"#,##0")&" Wh/day ' + ARROW + ' "&NominalVolts&" V","Chosen: "&NominalVolts&" V (auto rule suggests "&AutoVolts&" V)")&";  bank = "&Series&" ' + TIMES + ' "&BatteryUnitVolts&" V = "&TEXT(BankVolts,"0.0")&" V"')
prow("3. PV array", '=TEXT(RequiredWh,"#,##0.0")&" Wh ' + DIV + ' ("&PeakSunHours&" h ' + TIMES + ' "&SystemDerate&") = "&TEXT(PvRequiredW,"#,##0.0")&" W ' + ARROW + ' "&Panels&" ' + TIMES + ' "&PanelWatts&" W = "&TEXT(PvArrayW,"#,##0")&" W;  est. yield "&TEXT(DailyYieldWh,"#,##0")&" Wh/day"')
prow("4. Battery bank", '=TEXT(RequiredWh,"#,##0.0")&" Wh ' + TIMES + ' "&AutonomyDays&" days ' + DIV + ' ("&DepthOfDischarge&" ' + TIMES + ' "&BatteryEfficiency&") = "&TEXT(BatteryWhRequired,"#,##0.0")&" Wh ' + ARROW + ' "&TEXT(BatteryWhRequired/(BankVolts*BatteryUnitAh),"0.00")&" strings ' + ARROW + ' "&Parallel&" in parallel;  "&Series&"S"&Parallel&"P = "&TEXT(BankAh,"#,##0")&" Ah, usable "&TEXT(UsableWh,"#,##0")&" Wh"')
prow("5. Inverter", '=IF(AcLoadCount=0,"No AC loads: no inverter needed.","max("&TEXT(ContinuousW,"#,##0")&" W ' + TIMES + ' "&InverterSafetyFactor&" = "&TEXT(MinByContinuous,"#,##0.0")&" W;  surge "&TEXT(SurgeW,"#,##0")&" W ' + DIV + ' "&InverterSurgeMultiple&" = "&TEXT(MinBySurge,"#,##0.0")&" W) ' + ARROW + ' "&InverterW&" W (limited by "&LimitedBy&")")')
prow("6. MPPT controller", '=TEXT(PvArrayW,"#,##0")&" W ' + DIV + ' "&TEXT(BankVolts,"0.0")&" V = "&TEXT(ChargeCurrentA,"0.00")&" A ' + TIMES + ' "&ControllerSafetyFactor&" = "&TEXT(ControllerMinA,"0.0")&" A ' + ARROW + ' "&ControllerA&" A"')
prow("7. Battery-inverter cable", '=IF(ISNUMBER(InverterW),TEXT(InverterW,"#,##0")&" W ' + DIV + ' ("&TEXT(BankVolts,"0.0")&" V ' + TIMES + ' "&InverterEfficiency&") = "&TEXT(InverterInputA,"0.00")&" A;  fuse "&TEXT(InverterInputA*1.25,"0.0")&" A ' + ARROW + ' "&BI_FuseA&" A;  drop over "&CableBatteryInverterM&" m needs "&TEXT(BI_MinArea,"0.00")&" mm' + SQ + ';  ampacity >= fuse ' + ARROW + ' "&BI_Area&" mm' + SQ + ' ("&BI_Ampacity&" A, drop "&TEXT(BI_DropPct,"0.00")&" %)","n/a")')
prow("8. Controller-battery cable", '=TEXT(ChargeCurrentA,"0.00")&" A;  fuse "&TEXT(ChargeCurrentA*1.25,"0.0")&" A ' + ARROW + ' "&CB_FuseA&" A;  drop over "&CableControllerBatteryM&" m needs "&TEXT(CB_MinArea,"0.00")&" mm' + SQ + ';  ampacity >= fuse ' + ARROW + ' "&CB_Area&" mm' + SQ + ' ("&CB_Ampacity&" A, drop "&TEXT(CB_DropPct,"0.00")&" %)"')
prow("9. AC breaker", '=IF(ISNUMBER(InverterW),TEXT(InverterW,"#,##0")&" W ' + DIV + ' "&AcVolts&" V ' + TIMES + ' 1.25 = "&TEXT(AcBreakerMinA,"0.00")&" A ' + ARROW + ' "&AcBreakerA&" A MCB","n/a")')
r += 1
psection("Warnings")
prow("Count", '=IF(WarningCount=0,"None.",WarningCount)')
for i in range(len(warnings)):
    prow("", f"=INDEX(WarningList,{i + 1})")
r += 1
psection("Assumptions used")
prow("Site", '="Peak sun hours "&PeakSunHours&" h/day;  derate "&SystemDerate&";  autonomy "&AutonomyDays&" days;  max DC voltage drop "&MaxVoltageDropPct&" %"')
prow("Battery", '="Unit "&BatteryUnitVolts&" V / "&BatteryUnitAh&" Ah;  depth of discharge "&DepthOfDischarge&";  round-trip efficiency "&BatteryEfficiency')
prow("Inverter / controller", '="Efficiency "&InverterEfficiency&";  safety factor "&InverterSafetyFactor&";  surge multiple "&InverterSurgeMultiple&";  controller safety factor "&ControllerSafetyFactor&";  diversity "&DiversityFactor&";  AC "&AcVolts&" V"')
r += 1
wp.cell(row=r, column=1, value="Preliminary design aid, not a stamped engineering design. Check all results against local electrical codes, manufacturer datasheets and site conditions with a qualified person. Cable ampacities are indicative (single-core PVC copper, free air, ~30 C).").font = NOTE_FONT
wp.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
wp.cell(row=r, column=1).alignment = WRAP; wp.row_dimensions[r].height = 30
wp.column_dimensions["A"].width = 26; wp.column_dimensions["B"].width = 100; wp.column_dimensions["C"].width = 4
wp.print_area = f"A1:C{r}"
wp.page_setup.orientation = "portrait"; wp.page_setup.paperSize = wp.PAPERSIZE_A4
wp.page_setup.fitToWidth = 1; wp.page_setup.fitToHeight = 1
wp.sheet_properties.pageSetUpPr.fitToPage = True
wp.print_options.horizontalCentered = True
wp.page_margins.left = wp.page_margins.right = 0.5
wp.protection.sheet = True

# =====================================================================
# Read me
# =====================================================================
wm = wb.create_sheet("Read me", 3)
lines = [
    ("Off-Grid Solar Sizing Workbook", TITLE),
    ("A formula-driven Excel workbook for first-pass sizing of off-grid solar PV + battery systems. No macros. Excel 2016 or later (also works in LibreOffice Calc).", None),
    ("", None),
    ("How to use", BOLD),
    ("1. On the Inputs sheet, list your loads: appliance, watts, quantity, hours per day, start-up surge multiple (motors: 3-5; leave blank for 1) and whether the load is DC (wired to the battery, no inverter).", None),
    ("2. Check the site and equipment assumptions (yellow cells). Use worst-month peak sun hours for the site (PVGIS or NASA POWER). Set battery depth of discharge and efficiency for your chemistry.", None),
    ("3. Leave System voltage on 'auto' or force 12 / 24 / 48 V.", None),
    ("4. Read the Results sheet (every intermediate number) or print the Report sheet (one page, every formula step written out). File > Print on the Report sheet, or Save As PDF.", None),
    ("Yellow cells are inputs; grey cells are formulas. Sheets are protected without a password (Review > Unprotect Sheet) so formulas are not overwritten by accident.", None),
    ("", None),
    ("Method", BOLD),
    ("Daily energy = sum of DC Wh + sum of AC Wh / inverter efficiency.", None),
    ("System voltage (auto): 12 V up to 800 W continuous / 2 kWh per day, 24 V up to 2.5 kW / 5 kWh, otherwise 48 V.", None),
    ("PV array = daily Wh / (peak sun hours x derate), rounded up to whole panels.", None),
    ("Battery = daily Wh x autonomy days / (depth of discharge x round-trip efficiency); series count = system V / unit V; parallel strings rounded up.", None),
    ("Inverter = next standard size >= max(continuous AC load x safety factor, surge load / inverter surge multiple). Surge load = continuous + the largest single load's extra start-up power.", None),
    ("MPPT controller = array W / battery V x safety factor, next standard size.", None),
    ("Fuse = next standard size >= 1.25 x design current.", None),
    ("Cable = smallest copper size with voltage drop <= limit (2-wire DC run, out and back) AND ampacity >= fuse rating.", None),
    ("AC breaker = inverter W / AC voltage x 1.25, next standard MCB.", None),
    ("Standard sizes and the cable ampacity table are on the Tables sheet and can be edited (keep lists sorted ascending).", None),
    ("", None),
    ("Limits", BOLD),
    ("No PV string design: panel Voc at the coldest temperature is not checked against the controller's maximum input voltage.", None),
    ("No detailed protection coordination: fuse and breaker sizes are first estimates.", None),
    ("Cable ampacities are indicative for single-core PVC copper in free air at ~30 C; derate for conduit, heat and bundling.", None),
    ("Inverter surge capability is assumed to be the 'surge multiple' (default 2x) of its continuous rating; check the datasheet.", None),
    ("Up to 25 load rows. Lists on the Tables sheet are fixed length (add rows by unprotecting the sheet and extending the named ranges).", None),
    ("", None),
    ("Disclaimer", BOLD),
    ("This workbook is a preliminary design aid, not a stamped engineering design. Results must be checked against local electrical codes, manufacturer datasheets and site conditions by a qualified person. The author accepts no liability for designs based on it.", None),
    ("", None),
    ("License", BOLD),
    ("Licensed to the purchaser for personal and commercial use on their own projects, including client reports generated from it. Resale, redistribution or sharing of the workbook itself (modified or not) is not permitted.", None),
    ("", None),
    ("Verification", BOLD),
    ("The formulas mirror a tested JavaScript/Python engine. The clinic example on the Inputs sheet (1,695.7 Wh/day; 12 V; 1 x 550 W panel; 1S4P 400 Ah LiFePO4; 500 W surge-limited inverter; 60 A MPPT; 63 A fuses; 16 mm2 cables; 6 A MCB) was checked by automation in Excel.", None),
    ("", None),
    ("Author: Keshav Jha", NOTE_FONT),
]
for i, (text, font) in enumerate(lines, start=1):
    c = wm.cell(row=i, column=1, value=text); c.alignment = WRAP
    if font: c.font = font
wm.column_dimensions["A"].width = 120
wm.protection.sheet = True

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "names:", len(names))
